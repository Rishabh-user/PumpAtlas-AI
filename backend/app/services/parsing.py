"""Turn uploaded bytes and fetched pages into plain text for extraction.

Handles the formats vendor material actually arrives in: PDF datasheets, Word
questionnaires, Excel price lists and HTML product pages. Tables matter more than prose
in this domain, so the spreadsheet and PDF paths preserve row structure rather than
flattening everything into a paragraph.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from typing import Any

from app.core.logging import get_logger

log = get_logger(__name__)

_WS_RE = re.compile(r"[ \t\u00a0]+")
_BLANKS_RE = re.compile(r"\n{3,}")

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@dataclass
class ParsedContent:
    text: str
    page_count: int | None = None
    metadata: dict = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)
    content: dict = field(default_factory=dict)
    """The page as structure, not only as a wall of text.

    The flat ``text`` is what a reading model is given, and it is lossy in exactly the
    ways that matter here: a certification list stops being a list, a specification table
    stops being rows and columns, and a contact block stops being a contact. This keeps
    headings, paragraphs, tables, lists, links, documents and any address-like text, so a
    captured page can be re-read without being re-fetched and a person can check what was
    actually on it.
    """

    @property
    def char_count(self) -> int:
        return len(self.text)


#: Bumped whenever the parser starts keeping something it used to discard, so a source
#: captured under an older version can be recognised and re-parsed from its stored HTML
#: rather than re-fetched.
#:
#: 2 kept the page footer, which is where a company states its registered address,
#: company number and switchboard - the fields a vendor record most wants and a product
#: page never mentions. Before that, enriching a supplier from its own homepage came back
#: with a product family and no location.
PARSER_VERSION = 2

#: Caps for the structured capture. A source row is loaded on every screening pass, so an
#: unbounded blob is paid for again and again; these are generous enough to keep a long
#: "about us" page whole and mean enough to refuse a sitemap.
MAX_CAPTURE_HEADINGS = 80
MAX_CAPTURE_PARAGRAPHS = 200
MAX_CAPTURE_TABLES = 30
MAX_CAPTURE_TABLE_ROWS = 200
MAX_CAPTURE_LISTS = 60
MAX_CAPTURE_LIST_ITEMS = 60
MAX_CAPTURE_LINKS = 200
MAX_CAPTURE_CHARS = 4000

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]{2,}")

#: Deliberately conservative: an international prefix, or a bracketed area code followed
#: by a run of digits. A looser pattern turns every part number into a phone number.
PHONE_RE = re.compile(r"\+\d[\d\s().-]{7,}\d|\(\d{2,4}\)\s?\d[\d\s.-]{5,}\d")


def tidy(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _WS_RE.sub(" ", text)
    text = "\n".join(line.strip() for line in text.split("\n"))
    return _BLANKS_RE.sub("\n\n", text).strip()


def parse_pdf(data: bytes) -> ParsedContent:
    from pypdf import PdfReader

    warnings: list[str] = []
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:  # noqa: BLE001
            return ParsedContent(text="", warnings=["PDF is encrypted and could not be opened"])

    pages: list[str] = []
    for index, page in enumerate(reader.pages):
        try:
            pages.append(f"[page {index + 1}]\n{page.extract_text() or ''}")
        except Exception as exc:  # noqa: BLE001 - one bad page must not lose the rest
            warnings.append(f"page {index + 1} failed: {exc}")

    text = tidy("\n\n".join(pages))
    if len(text) < 200 and reader.pages:
        warnings.append(
            "Very little text recovered - this is likely a scanned document and needs OCR"
        )
    info = reader.metadata or {}
    return ParsedContent(
        text=text,
        page_count=len(reader.pages),
        metadata={
            "title": str(info.get("/Title") or "") or None,
            "author": str(info.get("/Author") or "") or None,
            "producer": str(info.get("/Producer") or "") or None,
        },
        warnings=warnings,
    )


def parse_docx(data: bytes) -> ParsedContent:
    import docx

    document = docx.Document(io.BytesIO(data))
    blocks = [p.text for p in document.paragraphs if p.text.strip()]
    for table_index, table in enumerate(document.tables):
        blocks.append(f"[table {table_index + 1}]")
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    props = document.core_properties
    return ParsedContent(
        text=tidy("\n".join(blocks)),
        metadata={"title": props.title or None, "author": props.author or None},
    )


def parse_xlsx(data: bytes, max_rows_per_sheet: int = 5000) -> ParsedContent:
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    sheet_names = [s.title for s in workbook.worksheets]
    blocks: list[str] = []
    warnings: list[str] = []
    for sheet in workbook.worksheets:
        blocks.append(f"[sheet: {sheet.title}]")
        for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
            if row_index >= max_rows_per_sheet:
                warnings.append(f"sheet {sheet.title} truncated at {max_rows_per_sheet} rows")
                break
            cells = ["" if value is None else str(value).strip() for value in row]
            if any(cells):
                blocks.append(" | ".join(cells))
    workbook.close()
    return ParsedContent(
        text=tidy("\n".join(blocks)),
        metadata={"sheets": sheet_names},
        warnings=warnings,
    )


def parse_csv(data: bytes) -> ParsedContent:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096])
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows = [" | ".join(cell.strip() for cell in row) for row in reader if any(row)]
    return ParsedContent(text=tidy("\n".join(rows)), metadata={"rows": len(rows)})


def parse_html(data: bytes | str) -> ParsedContent:
    from bs4 import BeautifulSoup

    html = data if isinstance(data, str) else _decode(data)
    soup = BeautifulSoup(html, "lxml")
    for tag in soup(["script", "style", "noscript", "nav", "svg"]):
        tag.decompose()

    # The footer stays. It was being decomposed with the scripts and the nav, and it is
    # where a company states the things a vendor record most wants and a product page
    # never says: registered office address, company registration number, switchboard
    # number, general enquiries address. Stripping it is why enriching a supplier from
    # its own homepage returned a product family and no location.
    #
    # Captured separately as well as left in the text, because "this came from the
    # footer" is useful to a reader judging a value.
    footer_text = "\n".join(
        text for tag in soup.find_all("footer") if (text := tag.get_text(" ", strip=True))
    )

    title = soup.title.string.strip() if soup.title and soup.title.string else None
    meta: dict[str, str | None] = {"title": title}
    for name, key in (
        ("description", "description"),
        ("author", "author"),
        ("article:published_time", "published_at"),
        ("og:site_name", "publisher"),
    ):
        tag = soup.find("meta", attrs={"name": name}) or soup.find("meta", attrs={"property": name})
        if tag and tag.get("content"):
            meta[key] = tag["content"].strip()

    capture: dict[str, Any] = {"title": title}

    capture["headings"] = [
        {"level": int(tag.name[1]), "text": text}
        for tag in soup.find_all(["h1", "h2", "h3", "h4"])
        if (text := tag.get_text(" ", strip=True))
    ][:MAX_CAPTURE_HEADINGS]

    # Lists before tables: certifications, approvals and product families are almost
    # always a <ul>, and flattening one into prose loses where each item ended.
    lists: list[list[str]] = []
    for element in soup.find_all(["ul", "ol"]):
        items = [
            text
            for item in element.find_all("li", recursive=False)
            if (text := item.get_text(" ", strip=True))
        ]
        if len(items) > 1:
            lists.append(items[:MAX_CAPTURE_LIST_ITEMS])
    capture["lists"] = lists[:MAX_CAPTURE_LISTS]

    blocks: list[str] = []
    if title:
        blocks.append(title)

    # Tables carry the datasheet values on most OEM product pages.
    tables: list[dict[str, Any]] = []
    for table_index, table in enumerate(soup.find_all("table")):
        blocks.append(f"[table {table_index + 1}]")
        rows: list[list[str]] = []
        for row in table.find_all("tr"):
            cells = [c.get_text(" ", strip=True) for c in row.find_all(["th", "td"])]
            if any(cells):
                rows.append(cells)
                blocks.append(" | ".join(cells))
        if rows and len(tables) < MAX_CAPTURE_TABLES:
            caption = table.find("caption")
            tables.append(
                {
                    "index": table_index + 1,
                    "caption": caption.get_text(" ", strip=True) if caption else None,
                    "rows": rows[:MAX_CAPTURE_TABLE_ROWS],
                }
            )
        table.decompose()
    capture["tables"] = tables

    capture["paragraphs"] = [
        text[:MAX_CAPTURE_CHARS]
        for tag in soup.find_all("p")
        if len(text := tag.get_text(" ", strip=True)) > 30
    ][:MAX_CAPTURE_PARAGRAPHS]

    body = soup.get_text("\n", strip=True)
    blocks.append(body)

    capture["links"] = [
        {"text": a.get_text(" ", strip=True)[:200], "href": a["href"]}
        for a in soup.find_all("a", href=True)
    ][:MAX_CAPTURE_LINKS]

    links = [
        a["href"]
        for a in soup.find_all("a", href=True)
        if a["href"].lower().endswith((".pdf", ".xlsx", ".docx"))
    ]
    meta["document_links"] = links[:50]
    capture["documents"] = links[:50]

    # Contact details, which a vendor record wants and prose loses.
    capture["emails"] = sorted({found.lower() for found in EMAIL_RE.findall(body)})[:20]
    capture["phones"] = sorted({found.strip() for found in PHONE_RE.findall(body)})[:20]
    capture["footer"] = footer_text[:MAX_CAPTURE_CHARS] or None
    capture["meta"] = {key: value for key, value in meta.items() if key != "title"}

    return ParsedContent(text=tidy("\n".join(blocks)), metadata=meta, content=capture)


def _decode(data: bytes) -> str:
    import chardet

    guess = chardet.detect(data[:20000])
    encoding = guess.get("encoding") or "utf-8"
    return data.decode(encoding, errors="replace")


EXTENSION_PARSERS = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".doc": parse_docx,
    ".xlsx": parse_xlsx,
    ".xlsm": parse_xlsx,
    ".csv": parse_csv,
    ".tsv": parse_csv,
    ".html": parse_html,
    ".htm": parse_html,
}

MIME_PARSERS = {
    PDF_MIME: parse_pdf,
    DOCX_MIME: parse_docx,
    XLSX_MIME: parse_xlsx,
    "text/csv": parse_csv,
    "text/html": parse_html,
    "application/xhtml+xml": parse_html,
}


def parse_bytes(
    data: bytes, filename: str | None = None, mime_type: str | None = None
) -> ParsedContent:
    """Dispatch on MIME type first, then extension, then fall back to plain text."""
    if mime_type:
        base = mime_type.split(";")[0].strip().lower()
        parser = MIME_PARSERS.get(base)
        if parser:
            return parser(data)
    if filename:
        suffix = filename.lower().rsplit(".", 1)
        if len(suffix) == 2:
            parser = EXTENSION_PARSERS.get(f".{suffix[1]}")
            if parser:
                return parser(data)
    if data[:5] == b"%PDF-":
        return parse_pdf(data)
    text = _decode(data)
    if "<html" in text[:2000].lower():
        return parse_html(text)
    return ParsedContent(text=tidy(text))
